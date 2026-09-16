"""Import the reviewed CV/GV rows from the local 405-point workbook.

The workbook is only the local third-stage evidence. Canonical locations were manually
checked one-by-one against KCMRIC first and WHO 2008 second before this importer was used.
"""
from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = Path(r"C:\Users\jiho3\OneDrive\바탕 화면\경혈학실습 연습용\경혈_데이터_405_WHO_해부학검수_재검수판.xlsx")

ENGLISH = {
    "CV": "Huiyin Qugu Zhongji Guanyuan Shimen Qihai Yinjiao Shenque Shuifen Xiawan Jianli Zhongwan Shangwan Juque Jiuwei Zhongting Shanzhong Yutang Zigong Huagai Xuanji Tiantu Lianquan Chengjiang".split(),
    "GV": "Changqiang Yaoshu Yaoyangguan Mingmen Xuanshu Jizhong Zhongshu Jinsuo Zhiyang Lingtai Shendao Shenzhu Taodao Dazhui Yamen Fengfu Naohu Qiangjian Houding Baihui Qianding Xinhui Shangxing Shenting Suliao Shuigou Duiduan Yinjiao".split(),
}

RESTRICTIONS = {
    "CV1": ("source_conflict", "KCMRIC은 직자법을 제시하지만 고전 문헌의 금침 기록도 병기함."),
    "CV2": ("conditional_contraindication", "임신부 자침 부적합/금기 기록."),
    "CV3": ("conditional_contraindication", "임신부는 신중히 자침하고 시술 전 배뇨."),
    "CV4": ("conditional_contraindication", "임신부 금침."),
    "CV5": ("conditional_contraindication", "여성 금침·금구라는 고전 기록이 있어 대상 조건 확인 필요."),
    "CV7": ("conditional_contraindication", "임신부 자침 주의."),
    "CV8": ("absolute_no_needling", "금침혈. 자침하지 않음."),
    "CV9": ("conditional_contraindication", "수병 환자 금침·구법 권고라는 고전 기록."),
    "CV14": ("depth_caution", "깊은 직자 및 위쪽 방향 자입 금지."),
    "CV15": ("source_conflict", "얕은 자법이 제시되나 고전 금침 기록도 있으며 깊은 직자·상향 사자 금지."),
    "CV16": ("depth_caution", "흉골 및 심장 방향의 깊은 직자 금지; 횡자 주의."),
    "CV17": ("source_conflict", "얕은 횡자법이 제시되나 고전 금침 기록도 병기됨."),
    "CV18": ("direction_caution", "흉부 정중선에서 깊은 직자 대신 얕은 횡자 방향 주의."),
    "CV19": ("direction_caution", "흉부 정중선에서 깊은 직자 대신 얕은 횡자 방향 주의."),
    "CV20": ("direction_caution", "흉부 정중선에서 깊은 직자 대신 얕은 횡자 방향 주의."),
    "CV21": ("direction_caution", "흉골 위 정중선에서 깊은 직자 대신 횡자 주의."),
    "CV22": ("direction_caution", "기관 손상 방지를 위해 직자 후 흉골 뒤를 따라 하향 전환하는 특수 자법; 깊은 직자 금지."),
    "CV23": ("direction_caution", "혀뿌리 쪽 상향 사자 시 설골·후두 구조 주의."),
    "GV1": ("depth_caution", "직장 손상 위험 때문에 깊은 자입 금지."),
    "GV15": ("depth_caution", "연수 손상 위험 때문에 위쪽으로 깊게 자입하지 않음."),
    "GV16": ("depth_caution", "연수 손상 위험 때문에 깊은 자입 금지."),
    "GV22": ("conditional_contraindication", "숫구멍이 닫히지 않은 소아(로컬 자료: 3세 이하)는 자침하지 않음."),
}


def parse_name(front: str) -> tuple[str, str]:
    match = re.search(r"^\s*([^\s(（]+)\s*[（(]([^）)]+)", front)
    if not match:
        raise ValueError(f"cannot parse name: {front!r}")
    return match.group(1), match.group(2)


def main() -> None:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    headers = [cell.value for cell in next(ws.iter_rows())]
    index = {name: headers.index(name) for name in headers if name is not None}
    rows = {"CV": [], "GV": []}
    for cells in ws.iter_rows(min_row=2, values_only=True):
        code = str(cells[index["WHO code"]] or "")
        meridian = code[:2]
        if meridian not in rows or not code[2:].isdigit():
            continue
        number = int(code[2:])
        korean, hanja = parse_name(str(cells[index["front"]]))
        needling_status, restriction = RESTRICTIONS.get(code, ("allowed_unverified", "문헌상 절대 금침 표기 없음; 자침 경로·깊이는 3D에서 미검증."))
        rows[meridian].append({
            "code": code,
            "front": cells[index["front"]],
            "korean": korean,
            "hanja": hanja,
            "english": ENGLISH[meridian][number - 1],
            "location": cells[index["location"]] or "",
            "method": cells[index["method"]] or "",
            "classification": cells[index["좌표화 분류"]] or "",
            "structures": cells[index["해부학적 구조물 기준\nEnglish (한글 구용어)"]] or "",
            "proportionRule": cells[index["비례촌/분할 좌표화 규칙"]] or "",
            "manualCheck": cells[index["수동 확인"]] or "",
            "whoStatus": cells[index["WHO 검수 상태"]] or "",
            "whoNote": cells[index["WHO 검수 내용"]] or "",
            "needlingStatus": needling_status,
            "needlingRestriction": restriction,
            "sourceValidation": {
                "order": ["KCMRIC", "WHO 2008", "local image"],
                "kcmric": "matched",
                "who": "matched",
                "localImage": "matched",
                "checkedOn": "2026-09-16",
            },
        })

    for meridian, points in rows.items():
        expected = 24 if meridian == "CV" else 28
        points.sort(key=lambda row: int(row["code"][2:]))
        if len(points) != expected:
            raise RuntimeError(f"{meridian}: expected {expected}, got {len(points)}")
        output = {
            "id": meridian,
            "meridian": "임맥" if meridian == "CV" else "독맥",
            "sourceFile": str(WORKBOOK),
            "sheet": ws.title,
            "note": "위치는 KCMRIC→WHO 2008→로컬 이미지 순서로 전수 1:1 검증. method는 로컬 실습 메모를 원문 보존하고 needlingRestriction은 KCMRIC 금기·주의를 별도 구조화함.",
            "points": points,
        }
        target = ROOT / "data" / "meridians" / "source" / f"{meridian}.json"
        target.write_text(json.dumps(output, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
        print(f"{target}: {len(points)} points")


if __name__ == "__main__":
    main()
